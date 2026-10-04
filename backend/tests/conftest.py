import io
import json
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
import pandas as pd
import fitz  # PyMuPDF
import docx

from app.main import app
from app.db.base import Base
from app.db.session import get_db
from app.core.config import settings

# In-memory SQLite for testing
SQLALCHEMY_DATABASE_URL = "sqlite:///:memory:"

engine = create_engine(
    SQLALCHEMY_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


@pytest.fixture(scope="function")
def db_session():
    Base.metadata.create_all(bind=engine)
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()
        Base.metadata.drop_all(bind=engine)


@pytest.fixture(scope="function")
def client(db_session):
    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


# --- Synthetic Document Generation Helpers ---

def create_synthetic_pdf(text: str = "Sample SOC Assessment Report\nDate: 2026-03-15\nIncident: INC-9901") -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text, fontsize=12)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_scanned_synthetic_pdf() -> bytes:
    """Creates a PDF with an image/shape but 0 text characters."""
    doc = fitz.open()
    page = doc.new_page()
    # Draw a rectangle simulating an image scan
    page.draw_rect(fitz.Rect(50, 50, 400, 400), color=(0.5, 0.5, 0.5), fill=(0.9, 0.9, 0.9))
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_synthetic_docx(title: str = "SOC Incident Closure Report", body: str = "Case ID: CASE-4029\nResolved by Tier 2.") -> bytes:
    doc = docx.Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(body)
    
    table = doc.add_table(rows=3, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Alert ID"
    hdr_cells[1].text = "Severity"
    hdr_cells[2].text = "Status"

    r1 = table.rows[1].cells
    r1[0].text = "ALT-101"
    r1[1].text = "High"
    r1[2].text = "Closed"

    r2 = table.rows[2].cells
    r2[0].text = "ALT-102"
    r2[1].text = "Medium"
    r2[2].text = "Investigating"

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def create_synthetic_csv(with_duplicates: bool = False, with_nulls: bool = False) -> bytes:
    rows = [
        {"alert_id": "ALT-001", "rule_name": "SSH Brute Force", "source_ip": "192.168.1.10", "severity": "High", "timestamp": "2026-02-10"},
        {"alert_id": "ALT-002", "rule_name": "Privilege Escalation", "source_ip": "10.0.0.50", "severity": "Critical", "timestamp": "2026-02-11"},
        {"alert_id": "ALT-003", "rule_name": "Port Scan", "source_ip": "172.16.0.22", "severity": "Low", "timestamp": "2026-02-12"},
    ]
    if with_duplicates:
        rows.append(rows[0])  # Duplicate
    if with_nulls:
        rows.append({"alert_id": "ALT-004", "rule_name": "", "source_ip": "", "severity": "Medium", "timestamp": "2026-02-13"})

    df = pd.DataFrame(rows)
    return df.to_csv(index=False).encode("utf-8")


def create_synthetic_xlsx() -> bytes:
    df_alerts = pd.DataFrame([
        {"alert_id": "ALT-100", "severity": "High", "mitre_technique": "T1059", "sla_met": "true"},
        {"alert_id": "ALT-101", "severity": "Critical", "mitre_technique": "T1078", "sla_met": "false"},
    ])
    df_metrics = pd.DataFrame([
        {"metric_name": "MTTR", "value": "45 mins", "target": "60 mins"},
        {"metric_name": "MTTD", "value": "12 mins", "target": "15 mins"},
    ])
    
    stream = io.BytesIO()
    with pd.ExcelWriter(stream, engine="openpyxl") as writer:
        df_alerts.to_excel(writer, sheet_name="Alerts", index=False)
        df_metrics.to_excel(writer, sheet_name="Performance", index=False)
    return stream.getvalue()


def create_synthetic_json() -> bytes:
    data = [
        {
            "case_id": "CASE-901",
            "closed_by": "Analyst_Alice",
            "resolution_code": "true_positive",
            "summary": "Phishing email contained malicious macro attachment.",
            "sla_hours": 3.5,
            "cve": "CVE-2024-21413",
        },
        {
            "case_id": "CASE-902",
            "closed_by": "Analyst_Bob",
            "resolution_code": "false_positive",
            "summary": "Internal penetration testing activity.",
            "sla_hours": 1.2,
            "cve": "CVE-2023-38606",
        }
    ]
    return json.dumps(data, indent=2).encode("utf-8")
