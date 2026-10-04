import io
from pathlib import Path
from typing import Union
import docx
from app.core.config import settings
from app.parsers.base import BaseParser, NormalizedDocument, ExtractedSection, ExtractedTable


class DocxParser(BaseParser):
    def parse(self, file_path_or_bytes: Union[str, Path, bytes], filename: str) -> NormalizedDocument:
        norm_doc = NormalizedDocument(
            document_type="docx",
            extraction_status="success",
        )
        try:
            if isinstance(file_path_or_bytes, bytes):
                file_stream = io.BytesIO(file_path_or_bytes)
                doc = docx.Document(file_stream)
            else:
                doc = docx.Document(str(file_path_or_bytes))

            text_parts = []
            sections = []
            current_section_title = "Introduction / Header"
            current_section_content = []

            for p_idx, p in enumerate(doc.paragraphs):
                p_text = p.text.strip()
                if not p_text:
                    continue

                style_name = p.style.name if p.style else ""
                
                # Check if paragraph is a heading
                is_heading = (
                    "heading" in style_name.lower()
                    or "title" in style_name.lower()
                    or (len(p_text) < 100 and p_text.isupper())
                )

                if is_heading:
                    # Flush previous section
                    if current_section_content:
                        sections.append(
                            ExtractedSection(
                                title=current_section_title,
                                level=1,
                                line_number=p_idx,
                                content="\n".join(current_section_content[:10]),
                            )
                        )
                        current_section_content = []
                    
                    current_section_title = p_text
                    level = 1
                    if "2" in style_name:
                        level = 2
                    elif "3" in style_name:
                        level = 3
                    
                    sections.append(
                        ExtractedSection(
                            title=p_text,
                            level=level,
                            line_number=p_idx,
                            content=p_text,
                        )
                    )
                else:
                    current_section_content.append(p_text)

                text_parts.append(p_text)

            # Flush last section
            if current_section_content:
                sections.append(
                    ExtractedSection(
                        title=current_section_title,
                        level=1,
                        content="\n".join(current_section_content[:10]),
                    )
                )

            # Extract tables
            extracted_tables = []
            for t_idx, table in enumerate(doc.tables):
                all_rows = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    all_rows.append(row_data)

                if all_rows:
                    headers = all_rows[0]
                    # Ensure unique and non-empty header names
                    clean_headers = []
                    seen = {}
                    for i, h in enumerate(headers):
                        h_name = h if h else f"Column_{i+1}"
                        if h_name in seen:
                            seen[h_name] += 1
                            h_name = f"{h_name}_{seen[h_name]}"
                        else:
                            seen[h_name] = 1
                        clean_headers.append(h_name)

                    sample_rows = []
                    for r in all_rows[1:settings.MAX_TABLE_PREVIEW_ROWS + 1]:
                        row_dict = {}
                        for i, h in enumerate(clean_headers):
                            row_dict[h] = r[i] if i < len(r) else ""
                        sample_rows.append(row_dict)

                    extracted_tables.append(
                        ExtractedTable(
                            name=f"DOCX Table {t_idx + 1}",
                            row_count=max(0, len(all_rows) - 1),
                            column_count=len(clean_headers),
                            headers=clean_headers,
                            rows=sample_rows,
                        )
                    )

            norm_doc.extracted_text = "\n\n".join(text_parts)
            norm_doc.sections = sections[:50]
            norm_doc.tables = extracted_tables

            if not norm_doc.extracted_text and not extracted_tables:
                norm_doc.extraction_status = "partial"
                norm_doc.warnings.append("DOCX file contains no readable text or tables.")

        except Exception as e:
            norm_doc.extraction_status = "failed"
            norm_doc.error_message = f"Failed to parse DOCX document: {str(e)}"
            norm_doc.warnings.append("DOCX parsing encountered an error.")

        return norm_doc
